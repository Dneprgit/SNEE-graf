from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


def set_run_font(run, size=24, bold=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle

    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    set_run_font(title_run, size=40, bold=True)

    subtitle_run = slide.placeholders[1].text_frame.paragraphs[0].runs[0]
    set_run_font(subtitle_run, size=22)


def add_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    set_run_font(title_run, size=30, bold=True)

    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()

    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
            p.text = text
            p.level = level
        else:
            p.text = item
            p.level = 0

        if p.runs:
            set_run_font(p.runs[0], size=19)


def main():
    output_path = (
        Path(__file__).resolve().parents[1]
        / "PRESENTATION_OPTIMIZATION_SIMPLE.pptx"
    )

    prs = Presentation()

    add_title_slide(
        prs,
        "Как работает расчет СНЭЭ",
        "Простое объяснение оптимальных параметров и графика заряда/разряда",
    )

    add_bullet_slide(
        prs,
        "Ключевой факт про оптимизацию",
        [
            "В названиях часто фигурирует QP (квадратичная оптимизация)",
            "Но в текущем коде используется линейная оптимизация (LP)",
            "Решатель: scipy linprog, метод HiGHS",
            "То есть фактическая математика сейчас линейная",
        ],
    )

    add_bullet_slide(
        prs,
        "Два шага расчета",
        [
            "Шаг 1: подбор оценочных параметров СНЭЭ",
            "Шаг 2: расчет почасового графика с выбранными параметрами",
            "Сначала определяем размер системы, потом ее режим работы",
        ],
    )

    add_bullet_slide(
        prs,
        "Что определяется на шаге 1",
        [
            "Подбираются сразу четыре величины:",
            ("входная мощность заряда (Nin)", 1),
            ("выходная мощность разряда (Nout)", 1),
            ("емкость батареи (C)", 1),
            ("максимальный остаточный дефицит (Dmax)", 1),
            "Одновременно учитывается почасовой баланс в 24 точках",
        ],
    )

    add_bullet_slide(
        prs,
        "Логика цели оптимизации",
        [
            "Главная цель - минимизировать дефицит",
            "Особенно сильно штрафуется пик дефицита (Dmax)",
            "Дополнительно ограничиваются слишком большие Nin, Nout и C",
            "Результат - компромисс между качеством покрытия и размером СНЭЭ",
        ],
    )

    add_bullet_slide(
        prs,
        "Как учитывается КПД",
        [
            "КПД связывает заряд и разряд через потери энергии",
            "Чтобы отдать 1 единицу в разряд, нужно зарядить больше 1 единицы",
            "Чем ниже КПД, тем выше потребность в заряде и мощности",
            "Из-за этого могут расти требуемые параметры системы",
        ],
    )

    add_bullet_slide(
        prs,
        "Что происходит на шаге 2",
        [
            "Берем найденные параметры (или заданные вручную)",
            "Считаем почасовой график заряда/разряда на 24 часа",
            "Формируем результирующий баланс и SOC батареи",
            "Получаем итоговый диспетчерский график для визуализации",
        ],
    )

    add_bullet_slide(
        prs,
        "Почему график именно такой",
        [
            "График - это результат решения задачи с ограничениями",
            "Заряд ограничен Nin, разряд ограничен Nout",
            "SOC всегда в пределах от 0 до емкости C",
            "Заряд и разряд ставятся туда, где это уменьшает дефицит лучше всего",
        ],
    )

    add_bullet_slide(
        prs,
        "Почему есть ровные участки и максимумы",
        [
            "Ровные участки появляются, когда выгодно держать стабильную мощность",
            "Пики на максимуме возникают, когда нужно покрыть сильный дефицит",
            "Если ограничений мало - график ближе к профилю",
            "Если ограничений много - график упирается в пределы",
        ],
    )

    add_bullet_slide(
        prs,
        "Принятые допущения модели",
        [
            "Горизонт: 24 часа, шаг: 1 час",
            "Параметры СНЭЭ постоянны в пределах суток",
            "Нет деградации батареи, тарифов и температурных эффектов",
            "Нет ограничений на скорость изменения мощности между часами",
            "Суточный цикл энергии замкнут (связь первого и последнего часа)",
        ],
    )

    add_bullet_slide(
        prs,
        "Как объяснять это пользователю",
        [
            "Сервис сначала подбирает размер СНЭЭ, потом считает режим работы",
            "КПД учитывает реальные потери, поэтому расчет не идеализирован",
            "График строится автоматически из ограничений и цели оптимизации",
            "Каждая точка графика имеет физическую и расчетную причину",
        ],
    )

    prs.save(str(output_path))
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
