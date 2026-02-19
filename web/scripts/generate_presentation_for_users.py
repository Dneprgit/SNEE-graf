from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


def set_run_font(run, size=24, bold=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold


def add_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title

    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    set_run_font(title_run, size=32, bold=True)

    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()

    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
            p.level = level
            p.text = text
        else:
            p.text = item
            p.level = 0

        if p.runs:
            set_run_font(p.runs[0], size=20)


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle

    title_run = slide.shapes.title.text_frame.paragraphs[0].runs[0]
    set_run_font(title_run, size=42, bold=True)

    subtitle_run = slide.placeholders[1].text_frame.paragraphs[0].runs[0]
    set_run_font(subtitle_run, size=24)


def main():
    output_path = Path(__file__).resolve().parents[1] / "PRESENTATION_FOR_USERS.pptx"

    prs = Presentation()

    add_title_slide(
        prs,
        "СНЭЭ Graf",
        "Понятная презентация сервиса для обычных пользователей",
    )

    add_bullet_slide(
        prs,
        "Что можно делать на сайте",
        [
            "Загружать суточный профиль баланса мощности (24 часа)",
            "Подбирать параметры СНЭЭ: мощность заряда, мощность разряда, емкость, КПД",
            "Интерактивно менять данные и сразу видеть результат",
            "Получать диспетчерский график, SOC и ключевые показатели эффективности",
        ],
    )

    add_bullet_slide(
        prs,
        "Какая исходная информация загружается",
        [
            "Профиль баланса мощности по часам (24 значения)",
            "Способы загрузки:",
            ("Excel-файл (.xlsx/.xls)", 1),
            ("Профиль по умолчанию", 1),
            ("Ручная правка значений в интерфейсе", 1),
            "Параметры СНЭЭ: Nin, Nout, емкость, КПД",
        ],
    )

    add_bullet_slide(
        prs,
        "Что рассчитывает система",
        [
            "Почасовой график заряда/разряда СНЭЭ",
            "Результирующий баланс после применения СНЭЭ",
            "Состояние заряда батареи (SOC) по каждому часу",
            "KPI: покрытие дефицита, использование избытка, пиковые мощности",
        ],
    )

    add_bullet_slide(
        prs,
        "Что можно интерактивно менять",
        [
            "Точки графика исходного профиля (drag-and-drop)",
            "Смещение всего профиля вверх/вниз",
            "Параметры батареи через поля ввода",
            "Параметры через SVG-батарею (перетаскивание маркеров)",
            "После изменения данные пересчитываются и графики обновляются",
        ],
    )

    add_bullet_slide(
        prs,
        "Как определяются оптимальные параметры СНЭЭ",
        [
            "Сначала запускается расчет оценочных параметров (QP-оптимизация)",
            "Одновременно подбираются:",
            ("входная мощность (Nin)", 1),
            ("выходная мощность (Nout)", 1),
            ("емкость батареи (C)", 1),
            ("остаточный дефицит мощности (Dmax)", 1),
            "Учитываются ограничения по КПД, мощности и емкости",
        ],
    )

    add_bullet_slide(
        prs,
        "Как строится диспетчерский график",
        [
            "Вход: профиль 24 часа + выбранные параметры СНЭЭ",
            "Оптимизатор рассчитывает почасовой режим заряд/разряд",
            "Для каждого часа формируется результирующий баланс",
            "Отдельно строится SOC-график с ограничением по емкости",
            "Пользователь видит эффект параметров в реальном времени",
        ],
    )

    add_bullet_slide(
        prs,
        "Как читать результат",
        [
            "Снижение дефицита: насколько СНЭЭ закрывает нехватку мощности/энергии",
            "Использование избытка: сколько излишней энергии удалось направить в заряд",
            "SOC: не выходит ли батарея за пределы допустимой емкости",
            "Сравнение исходного и результирующего баланса по каждому часу",
        ],
    )

    add_bullet_slide(
        prs,
        "Короткий сценарий демонстрации",
        [
            "1) Загрузить профиль (по умолчанию или Excel)",
            "2) Подправить 1-2 часа на графике",
            "3) Нажать: Рассчитать оценочные параметры",
            "4) Применить оценочные параметры",
            "5) Нажать: Рассчитать диспетчерский график",
            "6) Показать KPI и выгрузку в Excel",
        ],
    )

    add_bullet_slide(
        prs,
        "Польза для конечного пользователя",
        [
            "Быстрое понимание, какой накопитель нужен под конкретный профиль",
            "Прозрачная визуализация вместо ручных расчетов",
            "Возможность оперативно проверить разные сценарии",
            "Готовые показатели для принятия решения и отчета",
        ],
    )

    prs.save(str(output_path))
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
